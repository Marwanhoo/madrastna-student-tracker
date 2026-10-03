# -*- coding: utf-8 -*-
"""make_previews.py demo_raw.xlsx demo_recalculated.xlsx preview.xlsx

Turns the demo build into a static, values-only workbook that LibreOffice can render faithfully:
formulas become their computed values, ticks become box glyphs and the conditional formats are
painted as plain fills (LibreOffice draws conditional formats badly in PDF output).
Used only to produce the images in docs/ - the real workbook is never touched.
"""
import sys
import datetime as dt
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.formatting.formatting import ConditionalFormattingList

raw, calc, out = sys.argv[1:4]
wb, vals = load_workbook(raw), load_workbook(calc, data_only=True)
FIRST, CW, NCYC, LASTC = 7, 17, 12, 210
RED, ORANGE, GOLD, DARK = "FF8A8A", "FFB86B", "FFE08A", "3A3550"
KEEP = ["الرئيسية", "المواعيد", "مجموعة 1", "أونلاين"]
today = dt.date.today()


def paint(cell, color):
    cell.fill = PatternFill("solid", fgColor=color)


def as_date(v):
    return v.date() if isinstance(v, dt.datetime) else v


for name in list(wb.sheetnames):
    if name not in KEEP:
        del wb[name]
for ws in wb.worksheets:
    v = vals[ws.title]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.value = v[c.coordinate].value
    ws.conditional_formatting = ConditionalFormattingList()


def ticks(ws, head, r0, r1, absent):
    """box glyphs + the fills the conditional formats would give"""
    named = [r for r in range(r0, r1 + 1) if ws.cell(row=r, column=2).value]
    for c in range(FIRST, LASTC + 1):
        label = ws.cell(row=head, column=c).value
        held = any(ws.cell(row=r, column=c).value is True for r in range(r0, r1 + 1))
        for r in range(r0, r1 + 1):
            cell = ws.cell(row=r, column=c)
            on = cell.value is True
            if absent and label == "ح" and held and not on and r in named:
                paint(cell, RED)
            if label == "دفع" and not on and r in named and any(
                    ws.cell(row=r, column=x).value is True for x in range(c - CW + 1, c)):
                paint(cell, ORANGE)
    for c in range(FIRST, LASTC + 1):
        for r in range(r0, r1 + 1):
            cell = ws.cell(row=r, column=c)
            cell.value = "☑" if cell.value is True else "☐"
            cell.font = Font(name="DejaVu Sans", size=12, color=DARK)
            cell.alignment = Alignment(horizontal="center", vertical="center")


g = wb["مجموعة 1"]
ticks(g, 4, 5, 19, absent=True)
for c in range(FIRST, LASTC + 1):
    if as_date(g.cell(row=3, column=c).value) == today:
        paint(g.cell(row=3, column=c), GOLD)
for r in range(5, 20):
    d, e, f = (g.cell(row=r, column=x) for x in (4, 5, 6))
    if isinstance(d.value, (int, float)) and d.value >= 2:
        paint(d, RED)
    if isinstance(e.value, (int, float)) and e.value >= 2:
        paint(e, GOLD)
    if f.value not in (None, "", "تمام", "-"):
        paint(f, ORANGE)
for r in range(30, 45):
    rank = g.cell(row=r, column=5)
    if isinstance(rank.value, (int, float)) and rank.value <= 3:
        paint(rank, GOLD)
    for c in range(FIRST, LASTC + 1):
        s, m = g.cell(row=r, column=c), g.cell(row=28, column=c).value
        if isinstance(s.value, (int, float)) and isinstance(m, (int, float)):
            if s.value > m:
                paint(s, RED)
            elif s.value < m / 2:
                s.font = Font(name="Arial", size=10, bold=True, color="C0143C")

o = wb["أونلاين"]
ticks(o, 3, 4, 33, absent=False)
for r in range(4, 34):
    f = o.cell(row=r, column=6)
    if f.value == "آخر حصة":
        paint(f, GOLD)
    elif f.value == "تجديد" or str(f.value or "").startswith("عليه"):
        paint(f, ORANGE)

t = wb["المواعيد"]
for r in range(5, 35):
    if t.cell(row=r, column=9).value == 1:
        for c in range(1, 6):
            paint(t.cell(row=r, column=c), GOLD)

h = wb["الرئيسية"]
for r in range(18, 26):
    if str(h.cell(row=r, column=4).value or "").startswith("النهارده"):
        for c in range(2, 10):
            paint(h.cell(row=r, column=c), GOLD)

for ws, area, land in ((h, "A1:J56", False), (t, "A1:F12", False),
                       (g, "A1:AN25,A27:W37", True), (o, "A1:W12", True)):
    ws.print_area = area
    ws.page_setup.orientation = "landscape" if land else "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
wb.save(out)
print("preview workbook saved:", out)

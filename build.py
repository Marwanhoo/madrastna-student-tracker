# -*- coding: utf-8 -*-
"""Teacher student-management workbook (target: Google Sheets on a tablet).
Formulas only - no macros. Usage: python build.py out.xlsx [demo]"""
import sys
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1]
DEMO = len(sys.argv) > 2 and sys.argv[2] == "demo"

# ------------------------------------------------------------------ colours
PINK, YELLOW, MINT, SKY, LAV = "FF99C8", "FCF6BD", "D0F4DE", "A9DEF9", "E4C1F9"
CYCLE_COLORS = [PINK, YELLOW, MINT, SKY, LAV, "FEC89A", "D8E2DC", "FEC5BB",
                "C7CEEA", "FFD7BA", "E2F0CB", "E8E8E4"]
DARK, MUTED, WHITE = "3A3550", "8A85A0", "FFFFFF"
CALC = "F1F0F6"      # computed (do-not-edit) cells
LINE = "D5D1E0"
RED, ORANGE, GOLD = "FF8A8A", "FFB86B", "FFE08A"


def mix(hexcolor, t):
    """t=1 -> full colour, t=0 -> white"""
    r, g, b = (int(hexcolor[i:i + 2], 16) for i in (0, 2, 4))
    f = lambda v: round(255 - (255 - v) * t)
    return "%02X%02X%02X" % (f(r), f(g), f(b))


# ------------------------------------------------------------------ style helpers
_side = Side(style="thin", color=LINE)
BORDER = Border(left=_side, right=_side, top=_side, bottom=_side)
_fonts, _fills, _aligns = {}, {}, {}


def font(size=10, bold=False, color=DARK, italic=False):
    k = (size, bold, color, italic)
    if k not in _fonts:
        _fonts[k] = Font(name="Arial", size=size, bold=bold, color=color, italic=italic)
    return _fonts[k]


def fill(color):
    if color not in _fills:
        _fills[color] = PatternFill("solid", fgColor=color)
    return _fills[color]


def align(h="center", v="center", wrap=False):
    k = (h, v, wrap)
    if k not in _aligns:
        _aligns[k] = Alignment(horizontal=h, vertical=v, wrap_text=wrap, readingOrder=2)
    return _aligns[k]


def put(ws, row, col, value=None, bg=None, size=10, bold=False, color=DARK, h="center",
        wrap=False, border=True, fmt=None, italic=False, v="center"):
    c = ws.cell(row=row, column=col)
    if value is not None:
        c.value = value
    c.font = font(size, bold, color, italic)
    if bg:
        c.fill = fill(bg)
    c.alignment = align(h, v, wrap)
    if border:
        c.border = BORDER
    if fmt:
        c.number_format = fmt
    return c


def merge(ws, r1, c1, r2, c2, value=None, **kw):
    """style every cell of the block, write the anchor, then merge"""
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            put(ws, r, c, None, **kw)
    put(ws, r1, c1, value, **kw)
    if (r1, c1) != (r2, c2):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def width(ws, col, px):
    ws.column_dimensions[L(col)].width = round((px - 5) / 7, 2)


def height(ws, row, px):
    ws.row_dimensions[row].height = round(px * 0.75, 2)


def cf(ws, rng, formula, bg=None, fcolor=None, bold=None):
    kw = {}
    if bg:
        kw["fill"] = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
    if fcolor or bold is not None:
        kw["font"] = Font(color=fcolor or DARK, bold=bool(bold))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], **kw))


def setup_sheet(ws, tab_color, gridlines=False):
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = gridlines
    ws.sheet_properties.tabColor = tab_color


# ------------------------------------------------------------------ workbook
wb = Workbook()
wb.calculation.fullCalcOnLoad = True
wb._fonts = IndexedList([Font(name="Arial", size=10, family=2)])   # default font = Arial
wb._named_styles["Normal"].font = Font(name="Arial", size=10, family=2)

S_HOME, S_TIME, S_GUIDE, S_CALC = "الرئيسية", "المواعيد", "طريقة الاستخدام", "حسابات"
S_ONLINE, S_PRIV = "أونلاين", "برايفت"
NGROUPS = 8
GROUP_SHEETS = ["مجموعة %d" % i for i in range(1, NGROUPS + 1)]
IND_SHEETS = [S_ONLINE, S_PRIV]


def q(name):
    return "'%s'" % name


# shared grid geometry (group tabs and individual tabs use the same columns)
NCYC, NSES, CW = 12, 8, 17          # 12 "months" x 8 sessions; 17 columns per month
FIRST = 7                            # column G
LASTC = FIRST + NCYC * CW - 1        # column HB
G_, HB_ = L(FIRST), L(LASTC)
HC, HD, HE, HF, HG = (L(LASTC + i) for i in range(1, 6))   # hidden helper columns


def cyc_start(k):                    # first column of month k
    return FIRST + CW * (k - 1)


def col_h(k, j):                     # attendance column of session j in month k
    return cyc_start(k) + 2 * (j - 1)


def pay_col(k):                      # the "paid" column = last column of the month
    return cyc_start(k) + CW - 1


# helper-sheet addresses -------------------------------------------------------
DAYS = [("السبت", 7), ("الأحد", 1), ("الاثنين", 2), ("الثلاثاء", 3),
        ("الأربعاء", 4), ("الخميس", 5), ("الجمعة", 6)]          # WEEKDAY(date) numbers
TIMES = []
for hh in range(7, 24):
    for mm in (0, 30):
        h12 = hh if hh <= 12 else hh - 12
        TIMES.append("%d:%02d %s" % (h12, mm, "ص" if hh < 12 else "م"))
SEL = ["الكل"] + ["شهر %d" % k for k in range(1, NCYC + 1)]
C_ = q(S_CALC)
DAYNAMES = "%s!$A$2:$A$8" % C_
DAYNUMS = "%s!$B$2:$B$8" % C_
TIMELIST = "%s!$D$2:$D$%d" % (C_, 1 + len(TIMES))
SELLIST = "%s!$F$2:$F$%d" % (C_, 1 + len(SEL))
NTABS = NGROUPS + len(IND_SHEETS)
TABNAMES = "%s!$J$2:$J$%d" % (C_, 1 + NTABS)          # 8 groups + online + private
GRPNAMES = "%s!$J$2:$J$%d" % (C_, 1 + NGROUPS)
THR = "%s!$H$27" % q(S_HOME)                          # absence-alert threshold (input)

# timetable geometry
T0, TN = 5, 34                                        # 30 timetable rows
T_ = q(S_TIME)
TT_WHO = "%s!$D$%d:$D$%d" % (T_, T0, TN)
TT_TIME = "%s!$C$%d:$C$%d" % (T_, T0, TN)
TT_DNUM = "%s!$G$%d:$G$%d" % (T_, T0, TN)
TT_D = "%s!$K$%d:$K$%d" % (T_, T0, TN)
TT_D2 = "%s!$L$%d:$L$%d" % (T_, T0, TN)

GROUP_NST, IND_NST = 15, 30


# =============================================================== grid (shared)
def grid_columns(ws):
    width(ws, 1, 30)
    width(ws, 2, 140)
    width(ws, 3, 108)
    width(ws, 4, 46)
    width(ws, 5, 46)
    width(ws, 6, 80)
    for k in range(1, NCYC + 1):
        s = cyc_start(k)
        for c in range(s, s + CW - 1):          # the 16 session columns: collapsible group
            width(ws, c, 40)
            ws.column_dimensions[L(c)].outline_level = 1
        width(ws, pay_col(k), 104)               # "paid" column stays visible as the month summary
    for c in range(LASTC + 1, LASTC + 6):
        width(ws, c, 60)
        ws.column_dimensions[L(c)].hidden = True


def month_headers(ws, r_title, r_sess, label_fmt, r_date=None):
    """month title row, session-number row, optional date row"""
    for k in range(1, NCYC + 1):
        col = CYCLE_COLORS[k - 1]
        s = cyc_start(k)
        merge(ws, r_title, s, r_title, s + CW - 1, label_fmt % k, bg=col, size=12, bold=True)
        last = r_date or r_sess
        merge(ws, r_sess, pay_col(k), last, pay_col(k), "المصاريف", bg=mix(col, .8), size=9, bold=True)
        for j in range(1, NSES + 1):
            c = col_h(k, j)
            merge(ws, r_sess, c, r_sess, c + 1, "حصة %d" % j, bg=mix(col, .8), size=9, bold=True)
            if r_date:
                merge(ws, r_date, c, r_date, c + 1, None, bg=WHITE, size=10, bold=True, fmt="d/m")


def tick_block(ws, r_head, r0, r1):
    """sub-header + the tick rectangle (paid / attended / homework)"""
    for k in range(1, NCYC + 1):
        col = CYCLE_COLORS[k - 1]
        put(ws, r_head, pay_col(k), "دفع", bg=mix(col, .8), size=10, bold=True)
        for j in range(1, NSES + 1):
            c = col_h(k, j)
            put(ws, r_head, c, "ح", bg=mix(col, .55), size=10, bold=True)
            put(ws, r_head, c + 1, "و", bg=mix(col, .8), size=10, bold=True)
        for r in range(r0, r1 + 1):
            put(ws, r, pay_col(k), None, bg=mix(col, .7))
            for j in range(1, NSES + 1):
                c = col_h(k, j)
                put(ws, r, c, None, bg=mix(col, .18))
                put(ws, r, c + 1, None, bg=mix(col, .42))


def number_dv(ws):
    """score cells accept plain numbers only (a typed 8/10 would otherwise turn into a date)"""
    d = DataValidation(type="decimal", operator="between", formula1="0", formula2="1000", allow_blank=True,
                       showErrorMessage=True, errorTitle="الدرجة", error="اكتبي رقم بس، بالشكل ده: 8 أو 8.5")
    ws.add_data_validation(d)
    return d


def unpaid_cf(ws, r0, r1):
    for k in range(1, NCYC + 1):
        s = cyc_start(k)
        p = L(pay_col(k))
        cf(ws, "%s%d:%s%d" % (p, r0, p, r1),
           'AND($B%d<>"",N($%s%d)=0,SUMPRODUCT(--$%s%d:$%s%d)>0)'
           % (r0, p, r0, L(s), r0, L(s + CW - 2), r0), bg=ORANGE)


# =============================================================== group tab
def build_group(ws, default_name):
    setup_sheet(ws, SKY)
    grid_columns(ws)
    NST = GROUP_NST
    R_T, R_S, R_D, R_H = 1, 2, 3, 4
    ST0, ST1 = 5, 4 + NST                     # 5..19
    R_CNT = ST1 + 1                           # 20
    R_NL, R_NT, R_ML, R_MT = R_CNT + 2, R_CNT + 3, R_CNT + 4, R_CNT + 5      # 22..25
    R_BT, R_BM, R_BH = R_MT + 2, R_MT + 3, R_MT + 4                          # 27..29
    B0, B1 = R_BH + 1, R_BH + NST             # 30..44
    R_CYC, R_SES = B1 + 2, B1 + 3             # 46, 47 (hidden)
    rng = lambda r: "$%s$%d:$%s$%d" % (G_, r, HB_, r)
    row = lambda r: "$%s%d:$%s%d" % (G_, r, HB_, r)
    LAB, CNT, SES, CYC = rng(R_H), rng(R_CNT), rng(R_SES), rng(R_CYC)
    LABB, MAXR = rng(R_BH), rng(R_BM)

    for r, px in ((R_T, 32), (R_S, 22), (R_D, 26), (R_H, 22), (R_CNT, 22), (R_CNT + 1, 10),
                  (R_NL, 20), (R_NT, 56), (R_ML, 20), (R_MT, 30), (R_MT + 1, 14),
                  (R_BT, 28), (R_BM, 28), (R_BH, 22), (B1 + 1, 10)):
        height(ws, r, px)
    for r in list(range(ST0, ST1 + 1)) + list(range(B0, B1 + 1)):
        height(ws, r, 30)

    # ---- frozen corner
    merge(ws, R_T, 1, R_T, 2, default_name, bg=WHITE, size=12, bold=True)
    merge(ws, R_S, 1, R_D, 2,
          '=IF($%s$1=0,"لسه ما بدأتش",IF(INDEX(%s,1,17*($%s$2-1)+1)="","شهر "&$%s$2,'
          'INDEX(%s,1,17*($%s$2-1)+1))&" · حصة "&$%s$3&" من 8")'
          % (HC, rng(R_T), HC, HC, rng(R_T), HC, HC),
          bg=CALC, size=10, bold=True, color=MUTED)
    merge(ws, R_T, 3, R_D, 6,
          '="عدد الطلاب: "&$%s$1&CHAR(10)&"الغياب والواجب الناقص: في آخر 8 حصص"' % HE,
          bg=CALC, size=9, color=MUTED, wrap=True)
    month_headers(ws, R_T, R_S, "الشهر %d", r_date=R_D)

    # ---- attendance block
    for c, t in ((1, "م"), (2, "الاسم"), (3, "ولي الأمر"), (4, "غياب"), (5, "واجب"), (6, "المصاريف")):
        put(ws, R_H, c, t, bg=mix(LAV, .7), size=10, bold=True)
    tick_block(ws, R_H, ST0, ST1)
    absent = ('SUMPRODUCT((%s="ح")*(%s>0)*(%s<>TRUE)*(%s>$%s$1-8)*(%s>=$%s{r}))'
              % (LAB, CNT, "{row}", SES, HC, SES, HC))
    unpaid_terms = []
    for k in range(1, NCYC + 1):
        s = cyc_start(k)
        unpaid_terms.append("(COUNTIF($%s{r}:$%s{r},TRUE)>0)*($%s{r}<>TRUE)" % (L(s), L(s + CW - 2), L(pay_col(k))))
    for i, r in enumerate(range(ST0, ST1 + 1), 1):
        put(ws, r, 1, i, bg=CALC, size=9, color=MUTED)
        put(ws, r, 2, None, bg=WHITE, size=11, bold=True, h="right")
        put(ws, r, 3, None, bg=WHITE, size=10, fmt="@")
        ws["%s%d" % (HC, r)] = ('=97-SUMPRODUCT(MAX((%s="ح")*(%s=TRUE)*(97-%s)))' % (LAB, row(r), SES))
        ws["%s%d" % (HD, r)] = ('=SUMPRODUCT((%s="ح")*(%s=TRUE)*(%s>$%s$1-8))' % (LAB, row(r), SES, HC))
        ws["%s%d" % (HE, r)] = "=" + "+".join(t.format(r=r) for t in unpaid_terms)
        ws["%s%d" % (HF, r)] = ("=IF($%s%d=0,0,13-MAX(%s))" % (HE, r, ",".join(
            "%s*%d" % (t.format(r=r), 13 - k) for k, t in enumerate(unpaid_terms, 1))))
        put(ws, r, 4, '=IF($B{r}="","",IF(${hc}{r}=97,"-",IF(${hd}{r}=0,"متوقف",{a})))'
            .format(r=r, hc=HC, hd=HD, a=absent.format(r=r, row=row(r))), bg=CALC, size=10, bold=True)
        put(ws, r, 5, '=IF(OR($B{r}="",${hc}{r}=97),"",SUMPRODUCT(($G${h}:$HA${h}="ح")*($G{r}:$HA{r}=TRUE)'
            '*($H{r}:$HB{r}<>TRUE)*($G${s}:$HA${s}>${hc}$1-8)))'.format(r=r, hc=HC, h=R_H, s=R_SES),
            bg=CALC, size=10, bold=True)
        put(ws, r, 6, '=IF($B{r}="","",IF(${he}{r}=0,IF(${hc}{r}=97,"-","تمام"),IF(${he}{r}=1,'
            'IF(INDEX({t},1,17*(${hf}{r}-1)+1)="","شهر "&${hf}{r},INDEX({t},1,17*(${hf}{r}-1)+1)),'
            'IF(${he}{r}=2,"شهرين",${he}{r}&" شهور"))))'.format(r=r, he=HE, hf=HF, hc=HC, t=rng(R_T)),
            bg=CALC, size=9, bold=True)
    # counts row
    merge(ws, R_CNT, 1, R_CNT, 2, "العدد", bg=CALC, size=9, color=MUTED)
    merge(ws, R_CNT, 3, R_CNT, 6, None, bg=CALC)
    for c in range(FIRST, LASTC + 1):
        put(ws, R_CNT, c, "=COUNTIF(%s%d:%s%d,TRUE)" % (L(c), ST0, L(c), ST1),
            bg=CALC, size=9, color=MUTED, fmt="0;-0;;@")

    # ---- next-session box
    merge(ws, R_NL, 1, R_NL, 2, "الحصة الجاية: هنعمل إيه؟", bg=mix(MINT, .9), size=9, bold=True)
    merge(ws, R_NT, 1, R_NT, 2, None, bg=WHITE, size=10, wrap=True, h="right", fmt="@")
    merge(ws, R_ML, 1, R_ML, 2, "ميعاد مختلف؟ (اختياري)", bg=mix(MINT, .9), size=9, bold=True)
    merge(ws, R_MT, 1, R_MT, 2, None, bg=WHITE, size=10, wrap=True, h="right", fmt="@")
    merge(ws, R_NL, 3, R_NL, 6, "ميعادها حسب صفحة المواعيد", bg=mix(MINT, .9), size=9, bold=True)
    merge(ws, R_NT, 3, R_NT, 6, '=IF($%s$4="","مفيش ميعاد للمجموعة دي في صفحة المواعيد",$%s$4)' % (HF, HF),
          bg=CALC, size=10, bold=True, wrap=True)
    merge(ws, R_ML, 3, R_MT, 6, "لو الحصة اتأجلت أو في حصة إضافية، اكتبي ميعادها في الخانة البيضا",
          bg=CALC, size=8, color=MUTED, wrap=True)

    # ---- grades block
    merge(ws, R_BT, 1, R_BT, 2, "الدرجات · نوع الامتحان", bg=mix(LAV, .7), size=10, bold=True)
    merge(ws, R_BM, 1, R_BM, 2, "الدرجة النهائية (من كام)", bg=mix(LAV, .7), size=9, bold=True)
    merge(ws, R_BT, 3, R_BT, 5, "النسبة والترتيب على:", bg=mix(LAV, .7), size=9, bold=True)
    put(ws, R_BT, 6, "الكل", bg=WHITE, size=10, bold=True)
    merge(ws, R_BM, 3, R_BM, 6, "الدرجة ما بتتحسبش من غير نهائية", bg=CALC, size=8, color=MUTED)
    for c, t in ((1, "م"), (2, "الاسم"), (3, "المجموع"), (4, "%"), (5, "الترتيب"), (6, "التقدير")):
        put(ws, R_BH, c, t, bg=mix(LAV, .7), size=10, bold=True)
    for k in range(1, NCYC + 1):
        col = CYCLE_COLORS[k - 1]
        s = cyc_start(k)
        merge(ws, R_BT, pay_col(k), R_BM, pay_col(k), "ملاحظات" + CHR10 + "(سبب الغياب)", bg=mix(col, .8), size=8, bold=True, wrap=True)
        put(ws, R_BH, pay_col(k), "ملاحظات", bg=mix(col, .8), size=9, bold=True)
        for j in range(1, NSES + 1):
            c = col_h(k, j)
            merge(ws, R_BT, c, R_BT, c + 1, None, bg=WHITE, size=9, bold=True)
            put(ws, R_BM, c, None, bg=WHITE, size=10, bold=True)
            put(ws, R_BM, c + 1, None, bg=WHITE, size=10, bold=True)
            put(ws, R_BH, c, "كويز", bg=mix(col, .55), size=8, bold=True)
            put(ws, R_BH, c + 1, "امتحان", bg=mix(col, .8), size=8, bold=True)
        for r in range(B0, B1 + 1):
            put(ws, r, pay_col(k), None, bg=mix(col, .12), size=8, wrap=True, h="right", fmt="@")
            for j in range(1, NSES + 1):
                c = col_h(k, j)
                put(ws, r, c, None, bg=mix(col, .18), size=10)
                put(ws, r, c + 1, None, bg=mix(col, .42), size=10)
    for i, r in enumerate(range(B0, B1 + 1), 1):
        src = ST0 + i - 1
        put(ws, r, 1, i, bg=CALC, size=9, color=MUTED)
        put(ws, r, 2, '=IF($B%d="","",$B%d)' % (src, src), bg=CALC, size=11, bold=True, h="right")
        ws["%s%d" % (HC, r)] = ('=SUMIFS(%s,%s,">="&$%s$1,%s,"<="&$%s$2,%s,"<>ملاحظات",%s,">0")'
                                % (row(r), CYC, HD, CYC, HD, LABB, MAXR))
        ws["%s%d" % (HD, r)] = ('=SUMIFS(%s,%s,">="&$%s$1,%s,"<="&$%s$2,%s,"<>ملاحظات",%s,">=0")'
                                % (MAXR, CYC, HD, CYC, HD, LABB, row(r)))
        put(ws, r, 3, '=IF(OR($B{r}="",${hd}{r}=0),"",ROUND(${hc}{r},2)&" / "&${hd}{r})'.format(r=r, hc=HC, hd=HD),
            bg=CALC, size=10)
        put(ws, r, 4, '=IF(OR($B{r}="",${hd}{r}=0),"",${hc}{r}/${hd}{r})'.format(r=r, hc=HC, hd=HD),
            bg=CALC, size=10, bold=True, fmt="0%")
        put(ws, r, 5, '=IF($D{r}="","",RANK($D{r},$D${a}:$D${b}))'.format(r=r, a=B0, b=B1),
            bg=CALC, size=10, bold=True)
        put(ws, r, 6, '=IF($D{r}="","",IF($D{r}>=0.85,"ممتاز",IF($D{r}>=0.75,"جيد جداً",IF($D{r}>=0.65,"جيد",'
            'IF($D{r}>=0.5,"مقبول","ضعيف")))))'.format(r=r), bg=CALC, size=9, bold=True)

    # ---- hidden mask rows
    for k in range(1, NCYC + 1):
        s = cyc_start(k)
        for c in range(s, s + CW):
            ws.cell(row=R_CYC, column=c, value=k)
            ws.cell(row=R_SES, column=c, value=0)
        for j in range(1, NSES + 1):
            n = (k - 1) * NSES + j
            ws.cell(row=R_SES, column=col_h(k, j), value=n)
            ws.cell(row=R_SES, column=col_h(k, j) + 1, value=n)
    ws.row_dimensions[R_CYC].hidden = True
    ws.row_dimensions[R_SES].hidden = True

    # ---- hidden group-level helpers (rows 1-4 of the hidden columns)
    ws[HC + "1"] = '=SUMPRODUCT(MAX((%s="ح")*(%s>0)*%s))' % (LAB, CNT, SES)        # last session held
    ws[HC + "2"] = "=IF($%s$1=0,1,INT(($%s$1-1)/8)+1)" % (HC, HC)                 # current month
    ws[HC + "3"] = "=$%s$1-8*($%s$2-1)" % (HC, HC)                               # session inside the month
    ws[HD + "1"] = ('=IF(OR($F${r}="",$F${r}="الكل"),1,IFERROR(MATCH($F${r},{s},0)-1,1))'
                    .format(r=R_BT, s=SELLIST))
    ws[HD + "2"] = '=IF(OR($F${r}="",$F${r}="الكل"),{n},${hd}$1)'.format(r=R_BT, n=NCYC, hd=HD)
    ws[HE + "1"] = '=SUMPRODUCT(--($B$%d:$B$%d<>""))' % (ST0, ST1)               # students
    ws[HE + "2"] = "=IF(SUMPRODUCT((%s=TODAY())*(%s>0))>0,1,0)" % (rng(R_D), CNT)  # today's session recorded
    ws[HE + "3"] = ('=IF($A$1="",0,IF(${he}$2=1,SUMPRODUCT(MAX(({w}=$A$1)*(8-{d2}))),'
                    'SUMPRODUCT(MAX(({w}=$A$1)*(8-{d})))))'.format(he=HE, w=TT_WHO, d=TT_D, d2=TT_D2))
    ws[HF + "1"] = '=IF($%s$3=0,"",8-$%s$3)' % (HE, HE)                          # days until next session
    ws[HF + "2"] = '=IF($%s$1="","",TODAY()+$%s$1)' % (HF, HF)                   # its date
    ws[HF + "3"] = ('=IF(${hf}$1="",0,SUMPRODUCT(MAX(({w}=$A$1)*({n}=WEEKDAY(${hf}$2))*(ROW({w})-{o}))))'
                    .format(hf=HF, w=TT_WHO, n=TT_DNUM, o=T0 - 1))
    ws[HF + "4"] = ('=IF(${hf}$1="","",IF(${hf}$1=0,"النهارده",IF(${hf}$1=1,"بكرة",'
                    'INDEX({dn},MATCH(WEEKDAY(${hf}$2),{dw},0))))&" "&DAY(${hf}$2)&"/"&MONTH(${hf}$2)'
                    '&IF(${hf}$3=0,"",IF(INDEX({t},${hf}$3)="",""," – "&INDEX({t},${hf}$3))))'
                    .format(hf=HF, dn=DAYNAMES, dw=DAYNUMS, t=TT_TIME))
    ws[HG + "1"] = "=%s" % THR

    # ---- dropdowns
    dv = DataValidation(type="list", formula1='"وحدة,شهر"', allow_blank=True)
    dv.errorStyle, dv.showErrorMessage = "warning", True
    ws.add_data_validation(dv)
    for k in range(1, NCYC + 1):
        for j in range(1, NSES + 1):
            dv.add("%s%d" % (L(col_h(k, j)), R_BT))
    dv2 = DataValidation(type="list", formula1=SELLIST, allow_blank=False, showErrorMessage=True)
    ws.add_data_validation(dv2)
    dv2.add("F%d" % R_BT)
    dvn = number_dv(ws)
    dvd = DataValidation(type="date", operator="greaterThan", formula1="40000", allow_blank=True,
                         showErrorMessage=True, errorTitle="التاريخ", error="اكتبي التاريخ كده: 5/10")
    ws.add_data_validation(dvd)
    for k in range(1, NCYC + 1):
        a, b = L(cyc_start(k)), L(pay_col(k) - 1)
        dvn.add("%s%d:%s%d" % (a, R_BM, b, R_BM))
        dvn.add("%s%d:%s%d" % (a, B0, b, B1))
        dvd.add("%s%d:%s%d" % (a, R_D, b, R_D))

    # ---- conditional formats
    cf(ws, "%s%d:%s%d" % (G_, ST0, HB_, ST1),
       'AND($B%d<>"",%s$%d="ح",%s$%d>0,N(%s%d)=0)' % (ST0, G_, R_H, G_, R_CNT, G_, ST0), bg=RED)
    unpaid_cf(ws, ST0, ST1)
    cf(ws, "%s%d:%s%d" % (G_, R_D, HB_, R_D), 'AND(%s%d<>"",%s%d=TODAY())' % (G_, R_D, G_, R_D), bg=GOLD)
    cf(ws, "D%d:D%d" % (ST0, ST1), "AND(ISNUMBER(D%d),D%d>=$%s$1)" % (ST0, ST0, HG), bg=RED)
    cf(ws, "E%d:E%d" % (ST0, ST1), "AND(ISNUMBER(E%d),E%d>=2)" % (ST0, ST0), bg=GOLD)
    cf(ws, "F%d:F%d" % (ST0, ST1), 'AND(F%d<>"",F%d<>"تمام",F%d<>"-")' % (ST0, ST0, ST0), bg=ORANGE)
    sc = "%s%d:%s%d" % (G_, B0, HB_, B1)
    cf(ws, sc, 'AND(%s$%d<>"ملاحظات",ISNUMBER(%s%d),ISNUMBER(%s$%d),%s%d>%s$%d)'
       % (G_, R_BH, G_, B0, G_, R_BM, G_, B0, G_, R_BM), bg=RED)
    cf(ws, sc, 'AND(%s$%d<>"ملاحظات",ISNUMBER(%s%d),ISNUMBER(%s$%d),%s%d<%s$%d/2)'
       % (G_, R_BH, G_, B0, G_, R_BM, G_, B0, G_, R_BM), fcolor="C0143C", bold=True)
    cf(ws, "E%d:E%d" % (B0, B1), "AND(ISNUMBER(E%d),E%d<=3)" % (B0, B0), bg=GOLD)
    cf(ws, "%s%d:%s%d" % (G_, R_BM, HB_, R_BM),
       'AND(%s$%d<>"ملاحظات",NOT(ISNUMBER(%s%d)),COUNT(%s$%d:%s$%d)>0)'
       % (G_, R_BH, G_, R_BM, G_, B0, G_, B1), bg=ORANGE)

    ws.freeze_panes = "C4"
    return dict(ST0=ST0, ST1=ST1, R_NT=R_NT, R_MT=R_MT, R_D=R_D, R_BT=R_BT, R_BM=R_BM, B0=B0, B1=B1)


CHR10 = "\n"


# =============================================================== individual tab
def build_individual(ws, title, tab_color):
    setup_sheet(ws, tab_color)
    grid_columns(ws)
    NST = IND_NST
    R_T, R_S, R_H = 1, 2, 3
    ST0, ST1 = 4, 3 + NST                     # 4..33
    R_BT, R_BH = ST1 + 2, ST1 + 3             # 35, 36
    B0, B1 = R_BH + 1, R_BH + NST             # 37..66
    rng = lambda r: "$%s$%d:$%s$%d" % (G_, r, HB_, r)
    row = lambda r: "$%s%d:$%s%d" % (G_, r, HB_, r)
    LAB = rng(R_H)
    for r, px in ((R_T, 32), (R_S, 24), (R_H, 22), (ST1 + 1, 14), (R_BT, 28), (R_BH, 22)):
        height(ws, r, px)
    for r in list(range(ST0, ST1 + 1)) + list(range(B0, B1 + 1)):
        height(ws, r, 30)

    merge(ws, R_T, 1, R_T, 2, title, bg=WHITE, size=12, bold=True)
    merge(ws, R_S, 1, R_S, 2, '="عدد الطلاب: "&$%s$1' % HE, bg=CALC, size=10, bold=True, color=MUTED)
    merge(ws, R_T, 3, R_S, 6, "كل طالب بيتحسب لوحده: مع كل حصة علّمي المربع اللي عليه الدور",
          bg=CALC, size=9, color=MUTED, wrap=True)
    month_headers(ws, R_T, R_S, "الشهر %d (8 حصص)")
    for c, t in ((1, "م"), (2, "الاسم"), (3, "ولي الأمر"), (4, "الحصة"), (5, "باقي"), (6, "الحالة")):
        put(ws, R_H, c, t, bg=mix(LAV, .7), size=10, bold=True)
    tick_block(ws, R_H, ST0, ST1)
    for i, r in enumerate(range(ST0, ST1 + 1), 1):
        put(ws, r, 1, i, bg=CALC, size=9, color=MUTED)
        put(ws, r, 2, None, bg=WHITE, size=11, bold=True, h="right")
        put(ws, r, 3, None, bg=WHITE, size=10, fmt="@")
        ws["%s%d" % (HC, r)] = '=COUNTIFS(%s,"ح",%s,TRUE)' % (LAB, row(r))        # sessions taken
        ws["%s%d" % (HD, r)] = '=COUNTIFS(%s,"دفع",%s,TRUE)' % (LAB, row(r))      # months paid
        ws["%s%d" % (HE, r)] = "=8*$%s%d-$%s%d" % (HD, r, HC, r)                  # sessions left
        put(ws, r, 4, '=IF($B{r}="","",IF(${hc}{r}=0,"-",(${hc}{r}-8*INT((${hc}{r}-1)/8))&"/8"))'
            .format(r=r, hc=HC), bg=CALC, size=10, bold=True)
        put(ws, r, 5, '=IF($B{r}="","",IF(AND(${hc}{r}=0,${hd}{r}=0),"-",${he}{r}))'
            .format(r=r, hc=HC, hd=HD, he=HE), bg=CALC, size=10, bold=True)
        put(ws, r, 6, '=IF($B{r}="","",IF(AND(${hc}{r}=0,${hd}{r}=0),"جديد",IF(${he}{r}>1,"تمام",'
            'IF(${he}{r}=1,"آخر حصة",IF(${he}{r}=0,"تجديد",IF(${he}{r}=-1,"عليه حصة",IF(${he}{r}=-2,"عليه حصتين","عليه "&(-${he}{r})&" حصص")))))))'
            .format(r=r, hc=HC, hd=HD, he=HE), bg=CALC, size=9, bold=True)
    ws[HE + "1"] = '=SUMPRODUCT(--($B$%d:$B$%d<>""))' % (ST0, ST1)

    # grades block: score + "out of" per session
    merge(ws, R_BT, 1, R_BT, 2, "الدرجات", bg=mix(LAV, .7), size=10, bold=True)
    merge(ws, R_BT, 3, R_BT, 6, "اكتبي الدرجة، وجنبها النهائية (من)", bg=CALC, size=8, color=MUTED)
    for c, t in ((1, "م"), (2, "الاسم"), (3, "المجموع"), (4, "%"), (5, ""), (6, "التقدير")):
        put(ws, R_BH, c, t, bg=mix(LAV, .7), size=10, bold=True)
    for k in range(1, NCYC + 1):
        col = CYCLE_COLORS[k - 1]
        merge(ws, R_BT, pay_col(k), R_BT, pay_col(k), "ملاحظات", bg=mix(col, .8), size=9, bold=True)
        put(ws, R_BH, pay_col(k), "ملاحظات", bg=mix(col, .8), size=9, bold=True)
        for j in range(1, NSES + 1):
            c = col_h(k, j)
            merge(ws, R_BT, c, R_BT, c + 1, "حصة %d" % j, bg=mix(col, .8), size=9, bold=True)
            put(ws, R_BH, c, "درجة", bg=mix(col, .55), size=8, bold=True)
            put(ws, R_BH, c + 1, "من", bg=mix(col, .8), size=8, bold=True)
        for r in range(B0, B1 + 1):
            put(ws, r, pay_col(k), None, bg=mix(col, .12), size=8, wrap=True, h="right", fmt="@")
            for j in range(1, NSES + 1):
                c = col_h(k, j)
                put(ws, r, c, None, bg=mix(col, .18), size=10)
                put(ws, r, c + 1, None, bg=mix(col, .42), size=10)
    for i, r in enumerate(range(B0, B1 + 1), 1):
        src = ST0 + i - 1
        put(ws, r, 1, i, bg=CALC, size=9, color=MUTED)
        put(ws, r, 2, '=IF($B%d="","",$B%d)' % (src, src), bg=CALC, size=11, bold=True, h="right")
        ws["%s%d" % (HC, r)] = '=SUMIFS($G{r}:$HA{r},$G${h}:$HA${h},"درجة",$H{r}:$HB{r},">0")'.format(h=R_BH, r=r)
        ws["%s%d" % (HD, r)] = '=SUMIFS($H{r}:$HB{r},$G${h}:$HA${h},"درجة",$G{r}:$HA{r},">=0")'.format(h=R_BH, r=r)
        put(ws, r, 3, '=IF(OR($B{r}="",${hd}{r}=0),"",ROUND(${hc}{r},2)&" / "&${hd}{r})'.format(r=r, hc=HC, hd=HD),
            bg=CALC, size=10)
        put(ws, r, 4, '=IF(OR($B{r}="",${hd}{r}=0),"",${hc}{r}/${hd}{r})'.format(r=r, hc=HC, hd=HD),
            bg=CALC, size=10, bold=True, fmt="0%")
        put(ws, r, 5, None, bg=CALC)
        put(ws, r, 6, '=IF($D{r}="","",IF($D{r}>=0.85,"ممتاز",IF($D{r}>=0.75,"جيد جداً",IF($D{r}>=0.65,"جيد",'
            'IF($D{r}>=0.5,"مقبول","ضعيف")))))'.format(r=r), bg=CALC, size=9, bold=True)

    dvn = number_dv(ws)
    for k in range(1, NCYC + 1):
        dvn.add("%s%d:%s%d" % (L(cyc_start(k)), B0, L(pay_col(k) - 1), B1))
    cf(ws, "%s%d:%s%d" % (G_, B0, HB_, B1),
       'AND(%s$%d="من",ISNUMBER(F%d),NOT(ISNUMBER(%s%d)))' % (G_, R_BH, B0, G_, B0), bg=ORANGE)
    unpaid_cf(ws, ST0, ST1)
    cf(ws, "F%d:F%d" % (ST0, ST1), 'OR(F%d="تجديد",LEFT(F%d,4)="عليه")' % (ST0, ST0), bg=ORANGE)
    cf(ws, "F%d:F%d" % (ST0, ST1), 'F%d="آخر حصة"' % ST0, bg=GOLD)
    cf(ws, "%s%d:%s%d" % (G_, B0, HB_, B1),
       'AND(%s$%d="درجة",ISNUMBER(%s%d),ISNUMBER(%s%d),%s%d>%s%d)'
       % (G_, R_BH, G_, B0, L(FIRST + 1), B0, G_, B0, L(FIRST + 1), B0), bg=RED)
    ws.freeze_panes = "C3"
    return dict(ST0=ST0, ST1=ST1, B0=B0, B1=B1)


# =============================================================== build the sheets
ws_home = wb.active
ws_home.title = S_HOME
ws_time = wb.create_sheet(S_TIME)
G_INFO = [build_group(wb.create_sheet(n), n) for n in GROUP_SHEETS]
I_INFO = [build_individual(wb.create_sheet(S_ONLINE), S_ONLINE, MINT),
          build_individual(wb.create_sheet(S_PRIV), S_PRIV, LAV)]
ws_guide = wb.create_sheet(S_GUIDE)
ws_calc = wb.create_sheet(S_CALC)
GI, II = G_INFO[0], I_INFO[0]

# --------------------------------------------------------------- timetable
ws = ws_time
setup_sheet(ws, YELLOW)
for c, px in ((1, 30), (2, 100), (3, 100), (4, 175), (5, 170), (6, 12)):
    width(ws, c, px)
merge(ws, 1, 1, 1, 5, "جدول المواعيد الأسبوعي", bg=YELLOW, size=14, bold=True, border=False)
merge(ws, 2, 1, 2, 5, "صف لكل ميعاد ثابت. المجموعة اللي بتيجي مرتين في الأسبوع تتكتب في صفين.",
      bg=mix(YELLOW, .5), size=9, color=MUTED, border=False)
height(ws, 1, 40); height(ws, 2, 24); height(ws, 3, 8); height(ws, 4, 26)
for c, t in ((1, "م"), (2, "اليوم"), (3, "الساعة"), (4, "المجموعة / الطالب"), (5, "المكان / ملاحظات")):
    put(ws, 4, c, t, bg=mix(YELLOW, .9), size=10, bold=True)
for i, r in enumerate(range(T0, TN + 1), 1):
    height(ws, r, 30)
    put(ws, r, 1, i, bg=CALC, size=9, color=MUTED)
    put(ws, r, 2, None, bg=WHITE, size=11, bold=True)
    put(ws, r, 3, None, bg=WHITE, size=11)
    put(ws, r, 4, None, bg=WHITE, size=11, bold=True)
    put(ws, r, 5, None, bg=WHITE, size=10, wrap=True)
    ws["G%d" % r] = '=IF($B{r}="",0,IFERROR(INDEX({n},MATCH($B{r},{d},0)),0))'.format(r=r, n=DAYNUMS, d=DAYNAMES)
    ws["H%d" % r] = '=IF($C{r}="",99,IFERROR(MATCH($C{r},{t},0),99))'.format(r=r, t=TIMELIST)
    ws["I%d" % r] = "=IF(AND($G{r}>0,$G{r}=WEEKDAY(TODAY())),1,0)".format(r=r)
    ws["J%d" % r] = ('=IF($I{r}=1,COUNTIFS($I${a}:$I${b},1,$H${a}:$H${b},"<"&$H{r})'
                     '+COUNTIFS($I${a}:$I{r},1,$H${a}:$H{r},$H{r}),"")').format(r=r, a=T0, b=TN)
    ws["K%d" % r] = "=IF($G{r}=0,7,MOD($G{r}-WEEKDAY(TODAY()),7))".format(r=r)
    ws["L%d" % r] = "=IF($K{r}=0,7,$K{r})".format(r=r)
for c in "GHIJKL":
    ws.column_dimensions[c].hidden = True
dv = DataValidation(type="list", formula1=DAYNAMES, allow_blank=True, showErrorMessage=True)
dv.error, dv.errorTitle = "اختاري اليوم من القائمة", "اليوم"
ws.add_data_validation(dv); dv.add("B%d:B%d" % (T0, TN))
for rngx, src in (("C%d:C%d" % (T0, TN), TIMELIST), ("D%d:D%d" % (T0, TN), TABNAMES)):
    d = DataValidation(type="list", formula1=src, allow_blank=True)
    d.errorStyle, d.showErrorMessage = "warning", True
    ws.add_data_validation(d); d.add(rngx)
cf(ws, "A%d:E%d" % (T0, TN), "$I%d=1" % T0, bg=GOLD)
ws.freeze_panes = "A5"

# --------------------------------------------------------------- helper sheet
ws = ws_calc
ws["A1"], ws["B1"], ws["D1"], ws["F1"] = "اليوم", "WEEKDAY", "الساعة", "عرض الدرجات"
for i, (n, w) in enumerate(DAYS, 2):
    ws.cell(row=i, column=1, value=n); ws.cell(row=i, column=2, value=w)
for i, t in enumerate(TIMES, 2):
    ws.cell(row=i, column=4, value=t)
for i, t in enumerate(SEL, 2):
    ws.cell(row=i, column=6, value=t)
for c, t in zip("HIJKLMNOPQRSTU", ["#", "الصفحة", "الاسم", "طلاب", "آخر حصة", "الشهر", "حصة", "التقدم",
                                   "الحصة الجاية", "ميعاد يدوي", "ميعاد الجدول", "الميعاد", "نشطة", "تراكمي"]):
    ws[c + "1"] = t
for g, name in enumerate(GROUP_SHEETS, 1):
    r, s = g + 1, q(name)
    ws["H%d" % r], ws["I%d" % r] = g, name
    ws["J%d" % r] = '=IF({s}!$A$1="","{n}",{s}!$A$1)'.format(s=s, n=name)
    ws["K%d" % r] = "=%s!$%s$1" % (s, HE)
    ws["L%d" % r] = "=%s!$%s$1" % (s, HC)
    ws["M%d" % r] = "=%s!$%s$2" % (s, HC)
    ws["N%d" % r] = "=%s!$%s$3" % (s, HC)
    ws["O%d" % r] = "=%s!$A$2" % s
    ws["P%d" % r] = '=IF({s}!$A${t}="","",{s}!$A${t})'.format(s=s, t=GI["R_NT"])
    ws["Q%d" % r] = '=IF({s}!$A${t}="","",{s}!$A${t})'.format(s=s, t=GI["R_MT"])
    ws["R%d" % r] = "=%s!$%s$4" % (s, HF)
    ws["S%d" % r] = '=IF($Q{r}<>"",$Q{r},$R{r})'.format(r=r)
    ws["T%d" % r] = '=IF(OR($K{r}>0,$P{r}<>""),1,0)'.format(r=r)
    ws["U%d" % r] = "=$T%d" % r if g == 1 else "=$U%d+$T%d" % (r - 1, r)
for i, name in enumerate(IND_SHEETS, 1):
    r, s = NGROUPS + 1 + i, q(name)
    ws["H%d" % r], ws["I%d" % r] = NGROUPS + i, name
    ws["J%d" % r] = '=IF({s}!$A$1="","{n}",{s}!$A$1)'.format(s=s, n=name)
    ws["K%d" % r] = "=%s!$%s$1" % (s, HE)
for c, t in zip(["W", "X", "Y", "Z", "AA", "AB", "AC", "AD", "AE", "AF", "AG"],
                ["#", "المجموعة", "الاسم", "ولي الأمر", "غياب", "تنبيه غياب", "تراكمي", "المصاريف",
                 "تنبيه مصاريف", "تراكمي", "طالب"]):
    ws[c + "1"] = t
n = 0
for t, name in enumerate(GROUP_SHEETS + IND_SHEETS):
    is_group = t < NGROUPS
    info = GI if is_group else II
    s = q(name)
    for sr in range(info["ST0"], info["ST1"] + 1):
        n += 1
        r = n + 1
        ws["W%d" % r] = n
        ws["X%d" % r] = "=$J$%d" % (t + 2)
        ws["Y%d" % r] = '=IF({s}!$B{sr}="","",{s}!$B{sr})'.format(s=s, sr=sr)
        ws["Z%d" % r] = '=IF({s}!$C{sr}="","",{s}!$C{sr})'.format(s=s, sr=sr)
        if is_group:
            ws["AA%d" % r] = '=IF($Y{r}="","",{s}!$D{sr})'.format(r=r, s=s, sr=sr)
            ws["AD%d" % r] = ('=IF(OR($Y{r}="",{s}!$F{sr}="تمام",{s}!$F{sr}="-"),"",{s}!$F{sr})'
                              .format(r=r, s=s, sr=sr))
        else:
            ws["AA%d" % r] = '=""'
            ws["AD%d" % r] = ('=IF($Y{r}="","",IF(OR({s}!$F{sr}="تجديد",LEFT({s}!$F{sr},4)="عليه"),{s}!$F{sr},""))'
                              .format(r=r, s=s, sr=sr))
        ws["AB%d" % r] = '=IF(AND($Y{r}<>"",ISNUMBER($AA{r}),$AA{r}>={t}),1,0)'.format(r=r, t=THR)
        ws["AC%d" % r] = "=$AB%d" % r if n == 1 else "=$AC%d+$AB%d" % (r - 1, r)
        ws["AE%d" % r] = '=IF($AD{r}<>"",1,0)'.format(r=r)
        ws["AF%d" % r] = "=$AE%d" % r if n == 1 else "=$AF%d+$AE%d" % (r - 1, r)
        ws["AG%d" % r] = '=IF($Y{r}<>"",1,0)'.format(r=r)
NSTUD = n
SLAST = NSTUD + 1
ws.sheet_view.rightToLeft = True
ws.sheet_state = "hidden"
CS = lambda col: "%s!$%s$2:$%s$%d" % (C_, col, col, SLAST)      # student-table column
CG = lambda col: "%s!$%s$2:$%s$%d" % (C_, col, col, 1 + NGROUPS)  # group-table column

# --------------------------------------------------------------- dashboard
ws = ws_home
setup_sheet(ws, PINK)
width(ws, 1, 12)
for c in range(2, 10):
    width(ws, c, 71)
width(ws, 10, 12)
ws.column_dimensions["K"].hidden = True
B, I = 2, 9
merge(ws, 1, B, 1, I, "لوحة التحكم", bg=PINK, size=18, bold=True, border=False)
merge(ws, 2, B, 2, I, '="النهارده: "&INDEX({dn},MATCH(WEEKDAY(TODAY()),{dw},0))&" "&DAY(TODAY())&"/"'
      '&MONTH(TODAY())&"/"&YEAR(TODAY())'.format(dn=DAYNAMES, dw=DAYNUMS),
      bg=mix(PINK, .45), size=11, bold=True, border=False)
height(ws, 1, 46); height(ws, 2, 28); height(ws, 3, 10); height(ws, 4, 44); height(ws, 5, 24); height(ws, 6, 14)
kpis = [("=SUM(%s)" % CS("AG"), "طالب", PINK),
        ("=MAX(%s!$J$%d:$J$%d)" % (T_, T0, TN), "حصص النهارده", YELLOW),
        ("=%s!$AC$%d" % (C_, SLAST), "تنبيه غياب", MINT),
        ("=%s!$AF$%d" % (C_, SLAST), "مصاريف متأخرة", SKY)]
for i, (f, label, colr) in enumerate(kpis):
    c = B + 2 * i
    merge(ws, 4, c, 4, c + 1, f, bg=colr, size=20, bold=True, fmt="0")
    merge(ws, 5, c, 5, c + 1, label, bg=mix(colr, .55), size=10, bold=True)


def section(r, title, colr, c2=I):
    height(ws, r, 30)
    merge(ws, r, B, r, c2, title, bg=colr, size=12, bold=True, h="right")


def table_header(r, cols, colr):
    height(ws, r, 24)
    for c1, c2, t in cols:
        merge(ws, r, c1, r, c2, t, bg=mix(colr, .5), size=9, bold=True)


# today's sessions ------------------------------------------------
R = 7
section(R, "حصص النهارده", SKY)
table_header(R + 1, [(2, 2, "الساعة"), (3, 4, "المجموعة / الطالب"), (5, 6, "المكان"), (7, 9, "هنعمل إيه")], SKY)
for k in range(1, 7):
    r = R + 1 + k
    height(ws, r, 34)
    ws["K%d" % r] = '=IFERROR(MATCH(%d,%s!$J$%d:$J$%d,0),"")' % (k, T_, T0, TN)
    pick = lambda col: 'INDEX(%s!$%s$%d:$%s$%d,$K%d)' % (T_, col, T0, col, TN, r)
    merge(ws, r, 2, r, 2, '=IF($K{r}="","",{p})'.format(r=r, p=pick("C")), bg=WHITE, size=10, bold=True)
    empty = '"مفيش حصص متسجلة النهارده"' if k == 1 else '""'
    merge(ws, r, 3, r, 4, '=IF($K{r}="",{e},{p})'.format(r=r, e=empty, p=pick("D")),
          bg=WHITE, size=10, bold=True, wrap=True)
    merge(ws, r, 5, r, 6, '=IF($K{r}="","",IF({p}="","",{p}))'.format(r=r, p=pick("E")), bg=WHITE, size=9, wrap=True)
    merge(ws, r, 7, r, 9, '=IF($K{r}="","",IFERROR(INDEX({t},MATCH({p},{n},0)),""))'
          .format(r=r, t=CG("P"), p=pick("D"), n=GRPNAMES), bg=WHITE, size=9, wrap=True, h="right")

# next session per group -----------------------------------------
R = 16
section(R, "الحصة الجاية لكل مجموعة", MINT)
table_header(R + 1, [(2, 3, "المجموعة"), (4, 5, "الميعاد"), (6, 9, "هنعمل إيه")], MINT)
NEXT0 = R + 2
for k in range(1, NGROUPS + 1):
    r = R + 1 + k
    height(ws, r, 40)
    ws["K%d" % r] = '=IFERROR(MATCH(%d,%s,0),"")' % (k, CG("U"))
    merge(ws, r, 2, r, 3, '=IF($K{r}="","",INDEX({n},$K{r})&CHAR(10)&INDEX({o},$K{r}))'
          .format(r=r, n=CG("J"), o=CG("O")), bg=WHITE, size=9, bold=True, wrap=True)
    merge(ws, r, 4, r, 5, '=IF($K{r}="","",IF(INDEX({s},$K{r})="","—",INDEX({s},$K{r})))'
          .format(r=r, s=CG("S")), bg=WHITE, size=9, bold=True, wrap=True)
    merge(ws, r, 6, r, 9, '=IF($K{r}="","",INDEX({p},$K{r}))'.format(r=r, p=CG("P")),
          bg=WHITE, size=9, wrap=True, h="right")
cf(ws, "B%d:I%d" % (NEXT0, NEXT0 + NGROUPS - 1), 'LEFT($D%d,8)="النهارده"' % NEXT0, bg=GOLD)

# absence alerts ---------------------------------------------------
R = 27
section(R, "تنبيه الغياب (آخر 8 حصص)", mix(PINK, .75), c2=5)
merge(ws, R, 6, R, 7, "التنبيه من كام مرة؟", bg=mix(PINK, .75), size=9, bold=True)
merge(ws, R, 8, R, 9, 2, bg=WHITE, size=13, bold=True, fmt="0")
table_header(R + 1, [(2, 4, "الطالب"), (5, 5, "غياب"), (6, 7, "المجموعة"), (8, 9, "ولي الأمر")], PINK)
NABS = 10
for k in range(1, NABS + 1):
    r = R + 1 + k
    height(ws, r, 30)
    ws["K%d" % r] = '=IFERROR(MATCH(%d,%s,0),"")' % (k, CS("AC"))
    g = lambda col: '=IF($K{r}="","",INDEX({c},$K{r}))'.format(r=r, c=CS(col))
    merge(ws, r, 2, r, 4, g("Y"), bg=WHITE, size=10, bold=True, h="right")
    merge(ws, r, 5, r, 5, g("AA"), bg=WHITE, size=11, bold=True, color="C0143C")
    merge(ws, r, 6, r, 7, g("X"), bg=WHITE, size=9, wrap=True)
    merge(ws, r, 8, r, 9, g("Z"), bg=WHITE, size=10, fmt="@")
r = R + NABS + 2
height(ws, r, 22)
merge(ws, r, B, r, I, '=IF({t}>{n},"+ "&({t}-{n})&" كمان · عمود «غياب» في صفحات المجموعات",'
      'IF({t}=0,"مفيش تنبيهات غياب",""))'.format(t="%s!$AC$%d" % (C_, SLAST), n=NABS),
      size=9, color=MUTED, border=False)

# unpaid -----------------------------------------------------------
R = 41
section(R, "مصاريف متأخرة", "FEC89A")
table_header(R + 1, [(2, 4, "الطالب"), (5, 5, "المطلوب"), (6, 7, "المجموعة"), (8, 9, "ولي الأمر")], "FEC89A")
NPAY = 12
for k in range(1, NPAY + 1):
    r = R + 1 + k
    height(ws, r, 30)
    ws["K%d" % r] = '=IFERROR(MATCH(%d,%s,0),"")' % (k, CS("AF"))
    g = lambda col: '=IF($K{r}="","",INDEX({c},$K{r}))'.format(r=r, c=CS(col))
    merge(ws, r, 2, r, 4, g("Y"), bg=WHITE, size=10, bold=True, h="right")
    merge(ws, r, 5, r, 5, g("AD"), bg=WHITE, size=8, bold=True, color="B45309", wrap=True)
    merge(ws, r, 6, r, 7, g("X"), bg=WHITE, size=9, wrap=True)
    merge(ws, r, 8, r, 9, g("Z"), bg=WHITE, size=10, fmt="@")
r = R + NPAY + 2
height(ws, r, 22)
merge(ws, r, B, r, I, '=IF({t}>{n},"+ "&({t}-{n})&" كمان · عمود «المصاريف» في صفحات المجموعات",'
      'IF({t}=0,"كله دافع",""))'.format(t="%s!$AF$%d" % (C_, SLAST), n=NPAY),
      size=9, color=MUTED, border=False)
for rr in (15, 26, 40):
    height(ws, rr, 14)

# --------------------------------------------------------------- guide
ws = ws_guide
setup_sheet(ws, "E8E8E4")
width(ws, 1, 12); width(ws, 2, 560); width(ws, 3, 12)
tick_g = "%s%d:%s%d" % (G_, GI["ST0"], HB_, GI["ST1"])
tick_i = "%s%d:%s%d" % (G_, II["ST0"], HB_, II["ST1"])
GUIDE = [
    ("H", "طريقة الاستخدام", PINK),
    ("S", "القاعدة الأساسية", LAV),
    ("T", "• الخانات البيضا والملوّنة بلون الشهر: دي اللي بتكتبي أو بتعلّمي فيها."),
    ("T", "• الخانات الرمادي: بتتحسب لوحدها. ما تكتبيش فيها عشان المعادلات ما تتمسحش."),
    ("T", "• مثال للبيانات: الاسم «أحمد محمد علي» · ولي الأمر «01001234567». الصف ده مكتوب في مجموعة 1 "
          "وفي صفحة المواعيد كمثال، امسحيه واكتبي مكانه."),
    ("S", "في كل حصة (صفحة المجموعة)", SKY),
    ("T", "1. اكتبي اسم المجموعة في أول خانة فوق على اليمين (مرة واحدة)، والأسماء وأرقام أولياء الأمور."),
    ("T", "2. روحي لعمود الحصة الجديدة واكتبي تاريخها فوق، مثال: 5/10"),
    ("T", "3. علّمي تحت «ح» للي حضر، وتحت «و» للي عمل الواجب. اللي حضر زمايله وهو مش متعلّم بيتلوّن أحمر = غايب."),
    ("T", "4. لو الغياب بعذر: انزلي لجزء الدرجات واكتبي السبب في خانة «ملاحظات» بتاعة الشهر."),
    ("T", "5. تحت قايمة الأسماء: اكتبي «الحصة الجاية هنعمل إيه». ميعادها بيظهر لوحده من صفحة المواعيد."),
    ("S", "المصاريف", "FEC89A"),
    ("T", "• كل 8 حصص (شهر) ليهم لون وآخر عمود فيهم «دفع». علّمي عليه لما الطالب يدفع."),
    ("T", "• اللي حضر في الشهر ولسه ما دفعش خانته بتبقى برتقالي، واسمه بيظهر في الرئيسية."),
    ("T", "• الحصة الإضافية بتتسجل في العمود اللي عليه الدور عادي، وبتتحسب من الـ 8."),
    ("S", "الدرجات", MINT),
    ("T", "• تحت في نفس صفحة المجموعة. لكل حصة خانتين: «كويز» و«امتحان»."),
    ("T", "• اكتبي الدرجة النهائية في صف «الدرجة النهائية» الأول، وبعدها درجة كل طالب."),
    ("T", "• لو امتحان: اختاري نوعه (وحدة / شهر) من الخانة اللي فوقه."),
    ("T", "• المجموع والنسبة والترتيب والتقدير بيتحسبوا لوحدهم. من خانة «النسبة والترتيب على» اختاري «الكل» أو شهر معيّن."),
    ("S", "الأونلاين والبرايفت", LAV),
    ("T", "• كل طالب بيتحسب لوحده: مع كل حصة علّمي المربع اللي عليه الدور في صفّه، وعلّمي «دفع» لما يدفع الـ 8 حصص."),
    ("T", "• عمود «الحصة» بيقول هو في الحصة الكام من الـ 8، و«الحالة»: تمام / آخر حصة / تجديد / عليه حصص."),
    ("S", "المواعيد والرئيسية", YELLOW),
    ("T", "• في صفحة المواعيد: صف لكل ميعاد ثابت (اليوم، الساعة، المجموعة أو اسم الطالب)."),
    ("T", "• الرئيسية بتعرض: حصص النهارده، الحصة الجاية لكل مجموعة، اللي غاب كتير، واللي عليه مصاريف."),
    ("S", "خلّي بالك", "FEC5BB"),
    ("T", "• ما تعمليش ترتيب (Sort) ولا تضيفي أو تمسحي صفوف وأعمدة في صفحات المجموعات."),
    ("T", "• لو طالب مشي سيبي صفه زي ما هو، واكتبي الطالب الجديد في صف فاضي. الصف القديم علاماته بتفضل فيه."),
    ("T", "• اسم كل مجموعة لازم يكون مختلف. ولو غيّرتي اسم مجموعة، غيّريه كمان في صفحة المواعيد."),
    ("T", "• خانة «ميعاد مختلف» بتفضل ظاهرة في الرئيسية لحد ما تمسحيها."),
    ("T", "• «حصص النهارده» بتيجي من صفحة المواعيد بس. الحصة المتأجلة أو الإضافية بتظهر في جدول «الحصة الجاية»."),
    ("T", "• الأرقام تتكتب 123 مش ١٢٣. والدرجة رقم بس من غير «/»، والدرجة النهائية تتكتب في صفّها فوق."),
    ("T", "• الدرجة اللي من غير درجة نهائية ما بتتحسبش، وخانة النهائية بتاعتها بتبقى برتقالي."),
    ("T", "• لو خانة رمادي اتمسحت بالغلط: تراجع (Undo) على طول، أو من الكمبيوتر File > Version history."),
    ("S", "إزاي النظام بيحسب (افتراضات)", "E8E8E4"),
    ("T", "• السعة: 8 مجموعات × 15 طالب × 12 شهر (96 حصة). الأونلاين 30 طالب والبرايفت 30 طالب."),
    ("T", "• «غياب» = عدد مرات الغياب في آخر 8 حصص اتسجلت للمجموعة، من أول حصة الطالب حضرها."),
    ("T", "• الحصة بتتحسب «اتسجلت» لما طالب واحد على الأقل يتعلّم له حضور فيها."),
    ("T", "• الطالب اللي ما حضرش ولا حصة من آخر 8 بيتكتب قدامه «متوقف» وما بيظهرش في التنبيهات."),
    ("T", "• تنبيه الغياب من مرتين (الرقم ده بيتغيّر من الرئيسية). تلوين الواجب الناقص: من مرتين."),
    ("T", "• المصاريف: الطالب «عليه» شهر لو اتعلّم له أي حضور أو واجب فيه وخانة «دفع» فاضية."),
    ("T", "• النسبة = مجموع درجات الطالب ÷ مجموع الدرجات النهائية للخانات اللي ليه فيها درجة. "
          "الخانة الفاضية ما بتتحسبش؛ لو عايزاها صفر اكتبي 0."),
    ("T", "• التقدير: ممتاز 85% فأكتر · جيد جداً 75% · جيد 65% · مقبول 50% · أقل من كده ضعيف."),
    ("S", "تجهيز أول مرة (مرة واحدة، من متصفح على كمبيوتر)", PINK),
    ("T", "1. ارفع الملف على Google Drive وافتحه، وبعدين من القائمة: File > Save as Google Sheets"),
    ("T", "2. مربعات الصح: في كل صفحة مجموعة اكتب %s في خانة الاسم (Name box) واضغط Enter، وبعدين Insert > Checkbox. "
          "في «أونلاين» و«برايفت» النطاق %s" % (tick_g, tick_i)),
    ("T", "3. من File > Settings خلّي Locale على Egypt و Time zone على Cairo، عشان التاريخ 5/10 يتقري 5 أكتوبر و«النهارده» تبقى مظبوطة."),
    ("T", "4. الصفحات اللي مش محتاجاها (مجموعات زيادة): اعملها Hide من سهم الصفحة، وما تمسحهاش عشان الرئيسية بتقرأ منها."),
    ("T", "5. قبل ما تكتبوا أي بيانات: File > Make a copy واحتفظ بنسخة فاضية جاهزة للسنة الجاية."),
]
for i, item in enumerate(GUIDE, 1):
    if item[0] == "H":
        put(ws, i, 2, item[1], bg=item[2], size=16, bold=True, border=False); height(ws, i, 44)
    elif item[0] == "S":
        put(ws, i, 2, item[1], bg=mix(item[2], .8), size=11, bold=True, h="right", border=False); height(ws, i, 30)
    else:
        put(ws, i, 2, item[1], size=10, h="right", wrap=True, border=False)
        height(ws, i, 22 if len(item[1]) < 78 else 40)

# --------------------------------------------------------------- example rows
g1 = wb[GROUP_SHEETS[0]]
g1.cell(row=GI["ST0"], column=2, value="مثال: أحمد محمد علي")
g1.cell(row=GI["ST0"], column=3, value="01001234567")
ws_time["B5"], ws_time["C5"], ws_time["D5"], ws_time["E5"] = "السبت", "5:00 م", GROUP_SHEETS[0], "مثال، عدّليه"

# --------------------------------------------------------------- demo data (verification only)
if DEMO:
    today = dt.date.today()
    names = ["سارة أحمد", "يوسف خالد", "مريم علي", "عمر حسن", "ليلى محمود", "آدم سامي"]
    for i, nm in enumerate(names):
        g1.cell(row=5 + i, column=2, value=nm)
        g1.cell(row=5 + i, column=3, value="0100000000%d" % i)
    # 11 sessions held. attendance pattern per student (1 = present)
    att = {0: [1] * 11,                              # always present
           1: [1, 1, 1, 1, 0, 1, 0, 1, 1, 0, 1],     # 3 absences in the last 8 (sessions 4..11)
           2: [0, 0, 0, 0, 0, 1, 1, 0, 1, 1, 1],     # joined at session 6, 1 absence since
           3: [1, 1, 0, 0, 0, 0, 0, 0, 0, 0, 0],     # left after session 2 -> "متوقف"
           4: [1, 1, 1, 1, 1, 1, 1, 1, 1, 0, 0],     # 2 absences
           5: [1] * 11}
    for sidx, pat in att.items():
        for n_, a in enumerate(pat, 1):
            k, j = (n_ - 1) // 8 + 1, (n_ - 1) % 8 + 1
            if a:
                g1.cell(row=5 + sidx, column=col_h(k, j), value=True)
                if not (sidx == 5 and n_ in (9, 10)):              # student 6 misses homework twice
                    g1.cell(row=5 + sidx, column=col_h(k, j) + 1, value=True)
    for n_ in range(1, 12):
        k, j = (n_ - 1) // 8 + 1, (n_ - 1) % 8 + 1
        g1.cell(row=3, column=col_h(k, j), value=today - dt.timedelta(days=3 * (11 - n_)))
    for sidx in (0, 1, 3, 4, 5):
        g1.cell(row=5 + sidx, column=pay_col(1), value=True)      # month 1 paid
    for sidx in (0, 4):
        g1.cell(row=5 + sidx, column=pay_col(2), value=True)      # month 2 paid by two only
    g1.cell(row=5 + 2, column=pay_col(1), value=True)
    # grades: quiz in sessions 1,2 (month 1) and 9 (month 2); monthly exam in session 8
    g1.cell(row=GI["R_BM"], column=col_h(1, 1), value=10)
    g1.cell(row=GI["R_BM"], column=col_h(1, 2), value=10)
    g1.cell(row=GI["R_BM"], column=col_h(1, 8) + 1, value=50)
    g1.cell(row=GI["R_BT"], column=col_h(1, 8), value="شهر")
    g1.cell(row=GI["R_BM"], column=col_h(2, 1), value=20)
    sc = {0: (9, 10, 45, 18), 1: (7, 8, 30, 10), 2: (None, None, None, 15), 4: (5, 4, 20, 9), 5: (10, 9, 48, 20)}
    for sidx, (a, b_, c_, d_) in sc.items():
        r = GI["B0"] + sidx
        for v, colx in ((a, col_h(1, 1)), (b_, col_h(1, 2)), (c_, col_h(1, 8) + 1), (d_, col_h(2, 1))):
            if v is not None:
                g1.cell(row=r, column=colx, value=v)
    g1.cell(row=GI["B0"] + 1, column=pay_col(2), value="غاب 3 مرات - مريض")
    g1["A1"] = "علوم 5 ابتدائي"
    g1.cell(row=GI["R_NT"], column=1, value="مراجعة الوحدة الأولى + كويز")
    g2 = wb[GROUP_SHEETS[1]]
    g2["A1"] = "Math G6"
    g2["B5"], g2["C5"] = "Hana Tarek", "01111111111"
    g2.cell(row=GI["R_MT"], column=1, value="الخميس 7 م (حصة إضافية)")
    # timetable: two sessions today (unsorted), one tomorrow, one for group 2 in 3 days
    wd = {v: k for k, v in DAYS}
    dayname = lambda d: wd[(d.isoweekday() % 7) + 1]
    rows = [(dayname(today), "7:00 م", "Math G6", "أونلاين"),
            (dayname(today), "4:30 م", "علوم 5 ابتدائي", "السنتر"),
            (dayname(today + dt.timedelta(days=3)), "5:00 م", "علوم 5 ابتدائي", ""),
            (dayname(today + dt.timedelta(days=2)), "6:00 م", "كريم (برايفت)", "البيت")]
    for i, (d_, t_, w_, p_) in enumerate(rows):
        for c, v in zip("BCDE", (d_, t_, w_, p_)):
            ws_time["%s%d" % (c, T0 + i)] = v
    # individuals: taken / paid months
    on = wb[S_ONLINE]
    for i, (nm, taken, paid) in enumerate([("نور", 3, 1), ("زياد", 7, 1), ("جنى", 8, 1), ("مالك", 9, 1),
                                           ("فريدة", 3, 0), ("حمزة", 0, 0), ("ريم", 0, 1), ("تالا", 16, 2)]):
        r = II["ST0"] + i
        on.cell(row=r, column=2, value=nm)
        on.cell(row=r, column=3, value="0120000000%d" % i)
        for n_ in range(1, taken + 1):
            k, j = (n_ - 1) // 8 + 1, (n_ - 1) % 8 + 1
            on.cell(row=r, column=col_h(k, j), value=True)
        for k in range(1, paid + 1):
            on.cell(row=r, column=pay_col(k), value=True)
    on.cell(row=II["B0"], column=col_h(1, 1), value=8); on.cell(row=II["B0"], column=col_h(1, 1) + 1, value=10)
    on.cell(row=II["B0"], column=col_h(1, 2), value=15); on.cell(row=II["B0"], column=col_h(1, 2) + 1, value=20)

if DEMO:
    import os
    sel = os.environ.get("SEL")
    if sel:
        g1["F%d" % GI["R_BT"]] = sel
    def page(ws, area, landscape):
        ws.print_area = area
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
    from openpyxl.worksheet.properties import PageSetupProperties
    for w_ in wb.worksheets:
        w_.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    page(ws_home, "A1:J56", False)
    page(ws_time, "A1:F34", False)
    page(ws_guide, "A1:C45", False)
    for n_ in GROUP_SHEETS:
        page(wb[n_], "A1:AN45", True)
    for n_ in IND_SHEETS:
        page(wb[n_], "A1:AN45", True)
wb.save(OUT)
print("saved", OUT, "| students rows:", NSTUD, "| tick range group:", tick_g, "| individual:", tick_i)

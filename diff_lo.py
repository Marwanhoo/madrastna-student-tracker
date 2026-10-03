# -*- coding: utf-8 -*-
"""Compare the openpyxl original with the LibreOffice-recalculated file: formulas, CF, validations, structure."""
import sys, re
from openpyxl import load_workbook
raw, lo = sys.argv[1], sys.argv[2]
a, b = load_workbook(raw), load_workbook(lo)

def norm(f):
    if not isinstance(f, str):
        return f
    f = f.replace("TRUE()", "TRUE").replace("FALSE()", "FALSE")
    f = re.sub(r"'([^'!]+)'!", r"\1!", f)          # quotes around sheet names
    f = f.replace(" ", "")
    return f

diffs, n = [], 0
for name in a.sheetnames:
    wa, wb_ = a[name], b[name]
    for row in wa.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                n += 1
                v2 = wb_[c.coordinate].value
                if norm(v) != norm(v2):
                    diffs.append((name, c.coordinate, v[:120], str(v2)[:120]))
print("formulas compared:", n, "| differing:", len(diffs))
for d in diffs[:15]:
    print("  ", d)

def cfset(ws):
    out = []
    for rng, rules in ws.conditional_formatting._cf_rules.items():
        for r in rules:
            fill = r.dxf.fill.bgColor.rgb if (r.dxf and r.dxf.fill and r.dxf.fill.bgColor) else None
            font = r.dxf.font.color.rgb if (r.dxf and r.dxf.font and r.dxf.font.color) else None
            out.append((str(rng.sqref), norm(r.formula[0]), (fill or "")[-6:], (font or "")[-6:]))
    return sorted(out)
cfd = 0
for name in a.sheetnames:
    ca, cb = cfset(a[name]), cfset(b[name])
    if ca != cb:
        cfd += 1
        print("CF differs on", name, len(ca), len(cb))
        for x, y in zip(ca, cb):
            if x != y:
                print("   raw:", x); print("   lo :", y); break
print("sheets with CF differences:", cfd)

def dvset(ws):
    return sorted((str(d.sqref), d.type, norm(d.formula1), d.errorStyle or "stop", bool(d.showErrorMessage)) for d in ws.data_validations.dataValidation)
for name in a.sheetnames:
    da, db = dvset(a[name]), dvset(b[name])
    if da != db:
        print("DV differs on", name); print("   raw:", da[:2]); print("   lo :", db[:2])
print("DV check done")

for name in a.sheetnames:
    wa, wb_ = a[name], b[name]
    ma, mb = sorted(map(str, wa.merged_cells.ranges)), sorted(map(str, wb_.merged_cells.ranges))
    ol_a = sorted(k for k, d in wa.column_dimensions.items() if d.outline_level)
    lo_levels = {}
    for k, d in wb_.column_dimensions.items():
        if d.outline_level:
            lo_levels[k] = (d.min, d.max)
    n_lo = sum(mx - mn + 1 for mn, mx in lo_levels.values())
    print("%-16s rtl %s/%s | freeze %s/%s | merges %d/%d %s | state %s/%s | grouped cols %d/%d | tab %s/%s" % (
        name, wa.sheet_view.rightToLeft, wb_.sheet_view.rightToLeft, wa.freeze_panes, wb_.freeze_panes,
        len(ma), len(mb), "OK" if ma == mb else "DIFF", wa.sheet_state, wb_.sheet_state, len(ol_a), n_lo,
        wa.sheet_properties.tabColor.rgb[-6:] if wa.sheet_properties.tabColor else None,
        wb_.sheet_properties.tabColor.rgb[-6:] if wb_.sheet_properties.tabColor else None))
# fonts / number formats sample
g = b["مجموعة 1"]
print("fonts:", {g[c].font.name for c in ("A1", "B5", "G5", "D5", "G3", "G30")}, "| date fmt:", g["G3"].number_format, "| pct fmt:", g["D30"].number_format, "| phone fmt:", g["C5"].number_format, "| count fmt:", g["G20"].number_format)
print("row heights (5, 23, 30):", g.row_dimensions[5].height, g.row_dimensions[23].height, g.row_dimensions[30].height, "| col widths B,C,G,W:", g.column_dimensions["B"].width, g.column_dimensions["C"].width, g.column_dimensions["G"].width, g.column_dimensions["W"].width)

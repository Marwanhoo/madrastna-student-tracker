# -*- coding: utf-8 -*-
import sys
from openpyxl import load_workbook
f = sys.argv[1]
wb = load_workbook(f, data_only=True)
g1 = wb["مجموعة 1"]
print("== group 1 corner:", g1["A1"].value, "|", g1["A2"].value, "|", repr(g1["C1"].value))
print("L,c,p:", g1["HC1"].value, g1["HC2"].value, g1["HC3"].value, "| students:", g1["HE1"].value,
      "| doneToday:", g1["HE2"].value, "| m:", g1["HE3"].value, "| days:", g1["HF1"].value, "| date:", g1["HF2"].value, "| row:", g1["HF3"].value, "| txt:", g1["HF4"].value)
print("name | phone | absent | hw | fees | first | attWin | unpaidN | firstUnpaid")
for r in range(5, 12):
    print([g1.cell(row=r, column=c).value for c in (2, 3, 4, 5, 6)], [g1["%s%d" % (c, r)].value for c in ("HC", "HD", "HE", "HF")])
print("counts row (first 20 cols):", [g1.cell(row=20, column=c).value for c in range(7, 27)])
print("next box:", g1["A23"].value, "|", g1["C23"].value)
print("== grades (selector=%s) lo/hi: %s %s" % (g1["F27"].value, g1["HD1"].value, g1["HD2"].value))
for r in range(30, 37):
    print([g1.cell(row=r, column=c).value for c in (2, 3, 4, 5, 6)], g1["HC%d" % r].value, g1["HD%d" % r].value)
t = wb["المواعيد"]
print("== timetable")
for r in range(5, 10):
    print([t.cell(row=r, column=c).value for c in range(2, 6)], [t["%s%d" % (c, r)].value for c in "GHIJKL"])
on = wb["أونلاين"]
print("== online: name | session | left | status | taken | paid | rem")
for r in range(4, 13):
    print([on.cell(row=r, column=c).value for c in (2, 4, 5, 6)], [on["%s%d" % (c, r)].value for c in ("HC", "HD", "HE")])
print("online grades:", [on.cell(row=37, column=c).value for c in (2, 3, 4, 6)])
h = wb["الرئيسية"]
print("== dashboard")
for r in range(1, 57):
    vals = [(c.coordinate, c.value) for c in h[r] if c.value not in (None, "")]
    if vals:
        print(r, vals)
c = wb["حسابات"]
print("== calc groups")
for r in range(1, 12):
    print([c.cell(row=r, column=k).value for k in range(8, 22)])

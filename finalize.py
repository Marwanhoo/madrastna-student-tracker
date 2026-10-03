# -*- coding: utf-8 -*-
"""finalize.py raw.xlsx recalculated_copy.xlsx out.xlsx

LibreOffice's xlsx export squeezes fill colours into a 56-colour palette, which shifts the
teacher's chosen pastel palette. So the deliverable keeps the original (openpyxl) package and only
receives the *computed values* from the LibreOffice-recalculated copy, written as cached values.
"""
import sys
import zipfile
import datetime as dt
from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

raw, calc, out = sys.argv[1:4]
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PNS = "http://schemas.openxmlformats.org/package/2006/relationships"

values = load_workbook(calc, data_only=True)

zin = zipfile.ZipFile(raw)
wbx = etree.fromstring(zin.read("xl/workbook.xml"))
rels = etree.fromstring(zin.read("xl/_rels/workbook.xml.rels"))
rid_target = {r.get("Id"): r.get("Target") for r in rels.findall("{%s}Relationship" % PNS)}
part_sheet = {}
for sh in wbx.find("{%s}sheets" % NS):
    target = rid_target[sh.get("{%s}id" % RNS)].lstrip("/")
    if not target.startswith("xl/"):
        target = "xl/" + target
    part_sheet[target] = sh.get("name")

filled = {"number": 0, "text": 0, "empty": 0, "bool": 0}
zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename in part_sheet:
        ws = values[part_sheet[item.filename]]
        root = etree.fromstring(data)
        for c in root.iter("{%s}c" % NS):
            f = c.find("{%s}f" % NS)
            if f is None:
                continue
            for v in c.findall("{%s}v" % NS):
                c.remove(v)
            val = ws[c.get("r")].value
            v = etree.SubElement(c, "{%s}v" % NS)
            if isinstance(val, bool):
                c.set("t", "b"); v.text = "1" if val else "0"; filled["bool"] += 1
            elif isinstance(val, (dt.datetime, dt.date, dt.time)):
                c.attrib.pop("t", None); v.text = repr(float(to_excel(val))); filled["number"] += 1
            elif isinstance(val, (int, float)):
                c.attrib.pop("t", None)
                v.text = str(val) if isinstance(val, int) else repr(float(val)); filled["number"] += 1
            elif val is None:
                c.set("t", "str"); v.text = ""; filled["empty"] += 1
            else:
                if isinstance(val, str) and val.startswith("#"):
                    raise SystemExit("error value in %s!%s: %s" % (ws.title, c.get("r"), val))
                c.set("t", "str"); v.text = str(val); filled["text"] += 1
        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    zout.writestr(item, data)
zout.close()
print("cached values written:", filled)
